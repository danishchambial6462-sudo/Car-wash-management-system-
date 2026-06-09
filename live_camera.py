import cv2
import easyocr
from ultralytics import YOLO
from datetime import datetime
from collections import Counter
import time
import os
from openpyxl import Workbook, load_workbook

reader = easyocr.Reader(['en'])
plate_model = YOLO("license_plate_detector.pt")
def fix_plate(text):

    text = text.upper().replace(" ", "")

    if len(text) < 10:
        return text

    text = list(text[:10])

    for i in [2, 3, 6, 7, 8, 9]:

        if text[i] == "O":
            text[i] = "0"

        if text[i] == "I":
            text[i] = "1"

        if text[i] == "Z":
            text[i] = "2"

        if text[i] == "S":
            text[i] = "5"

        if text[i] == "E":
            text[i] = "6"

    return "".join(text)

cap = cv2.VideoCapture(0)

print("Press SPACE to capture")
print("Press Q to quit")

saved_numbers = set()

while True:

    ret, frame = cap.read()

    if not ret:
        break

    cv2.imshow("Car Wash Camera", frame)

    key = cv2.waitKey(1)

    if key == 32:

        print("\nCapturing... Hold Camera Steady")

        readings = []

        for i in range(3):

            ret, temp_frame = cap.read()

            if not ret:
                continue
            # YOLO Plate Detection

            results = plate_model(temp_frame)

            plate_found = False

            for r in results:

                 for box in r.boxes:

                   x1, y1, x2, y2 = map(int, box.xyxy[0])

                   plate_crop = temp_frame[y1:y2, x1:x2]

                   if plate_crop.size > 0:

                       temp_frame = plate_crop
                       plate_found = True
                       break

                 if plate_found:
                    break

            if not plate_found:

                 print("No Plate Detected")
                 continue 

            # Image Preprocessing
            temp_frame = cv2.resize(temp_frame, None, fx=3, fy=3)
            gray = cv2.cvtColor(temp_frame, cv2.COLOR_BGR2GRAY)

            gray = cv2.GaussianBlur(gray, (3, 3), 0)

            _, thresh = cv2.threshold(
                gray,
                0,
                255,
                cv2.THRESH_BINARY + cv2.THRESH_OTSU
            )

            result = reader.readtext(
                  thresh,
                  allowlist='ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789'
            )

            for item in result:

                text = item[1].upper()
                text = fix_plate(text)

                text = text.replace(" ", "")
                text = text.replace("[", "")
                text = text.replace("]", "")
                text = text.replace("'", "")
                text = text.replace('"', "")
                text = text.replace("|", "")

                print("OCR RAW:", text)

                if 8 <= len(text) <=10:
                    readings.append(text)

            time.sleep(0.5)

        print("All Readings:", readings)

        if len(readings) > 0:

            vehicle_no = Counter(readings).most_common(1)[0][0]

            print("FINAL VEHICLE NUMBER:", vehicle_no)

            if vehicle_no not in saved_numbers:

                saved_numbers.add(vehicle_no)

                os.makedirs("captures", exist_ok=True)

                current_time = datetime.now()

                photo_name = (
                    f"captures/{vehicle_no}_"
                    f"{current_time.strftime('%Y%m%d_%H%M%S')}.jpg"
                )

                cv2.imwrite(photo_name, frame)

                excel_file = "vehicles.xlsx"

                if not os.path.exists(excel_file):

                    wb = Workbook()
                    ws = wb.active

                    ws.append([
                        "Vehicle Number",
                        "Date",
                        "Time",
                        "Diesel Wash",
                        "Shinning Spray",
                        "Amount",
                        "Photo"
                    ])

                    wb.save(excel_file)

                wb = load_workbook(excel_file)
                ws = wb.active

                diesel = input("Diesel Wash? (y/n): ").lower() == "y"
                spray = input("Shining Spray? (y/n): ").lower() =="y"

                amount = 300
                if diesel:
                    amount +=50

                if spray:
                    amount += 100

                ws.append([
                    vehicle_no,
                    current_time.strftime("%d-%m-%y"),
                    current_time.strftime("%I:%M:%S %p"),
                    "Yes" if diesel else "No",
                    "yes" if spray else "No",
                    amount,
                    photo_name
                ])

                wb.save(excel_file)

                print("Date:", current_time.strftime("%d-%m-%y"))
                print("Time:", current_time.strftime("%I:%M:%S %p"))
                print("Photo Saved:", photo_name)
                print("Saved Successfully!")

            else:

                print("Duplicate Vehicle - Not Saved")

        else:

            print("Number Plate Not Found")

    elif key == ord('q'):
        break

cap.release()
cv2.destroyAllWindows()